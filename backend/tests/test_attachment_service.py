"""
Tests unitaires — Module 2 (extension) : Pièces jointes
(app/services/attachment_service.py)

Extraction testée avec de vrais fichiers construits en mémoire (docx/xlsx/
pptx via leurs libs respectives ; PDF construit octet par octet — pypdf
n'a pas d'API d'écriture de texte, cf. exploration manuelle avant d'écrire
ces tests). Vision testée avec un provider factice déterministe (pas de
dépendance à un vrai Ollama, cf. tests/test_vision_layer.py pour le
provider réel).

Toutes les données sont fictives.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.agents.vision_layer import BaseVisionProvider, VisionProviderError
from app.models.session import AttachmentModel
from app.services import attachment_service


def _build_minimal_pdf(text: str) -> bytes:
    """Construit un PDF minimal valide (1 page, police Helvetica) avec
    `text` comme contenu — pypdf n'offrant pas d'API d'écriture de texte,
    cf. exploration manuelle (xref/trailer construits à la main, validé en
    relisant avec PdfReader avant d'écrire ce test)."""
    content_stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objs = [
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n",
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n",
        b"3 0 obj<</Type/Page/Parent 2 0 R/Resources<</Font<</F1 4 0 R>>>>"
        b"/MediaBox[0 0 200 200]/Contents 5 0 R>>endobj\n",
        b"4 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n",
        b"5 0 obj<</Length " + str(len(content_stream)).encode() + b">>\nstream\n"
        + content_stream + b"\nendstream\nendobj\n",
    ]
    body = b"%PDF-1.4\n"
    offsets = [0]
    for o in objs:
        offsets.append(len(body))
        body += o
    xref_offset = len(body)
    n = len(objs) + 1
    xref = f"xref\n0 {n}\n0000000000 65535 f \n".encode()
    for off in offsets[1:]:
        xref += f"{off:010d} 00000 n \n".encode()
    trailer = f"trailer<</Size {n}/Root 1 0 R>>\nstartxref\n{xref_offset}\n%%EOF".encode()
    return body + xref + trailer


def _build_docx(text: str) -> bytes:
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_xlsx(cell_value: str) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = cell_value
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_pptx(text: str) -> bytes:
    from io import BytesIO

    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height // 4)
    box.text_frame.text = text
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


class _FixedVisionProvider(BaseVisionProvider):
    @property
    def provider_name(self) -> str:
        return "fixed"

    async def describe(self, image_bytes: bytes) -> str:
        return "[FICTIF] Description factice de l'image."


class _FailingVisionProvider(BaseVisionProvider):
    @property
    def provider_name(self) -> str:
        return "failing"

    async def describe(self, image_bytes: bytes) -> str:
        raise VisionProviderError("modèle vision indisponible", provider="failing")


class TestClassifyAttachment:
    def test_image_content_type(self):
        assert attachment_service.classify_attachment("photo.png", "image/png") == "image"

    def test_document_extension(self):
        assert attachment_service.classify_attachment("spec.pdf", "application/pdf") == "document"

    def test_generic_content_type_falls_back_to_extension(self):
        assert (
            attachment_service.classify_attachment("notes.txt", "application/octet-stream") == "document"
        )


class TestSaveAttachment:
    @pytest.mark.asyncio
    async def test_writes_file_under_session_dir_named_by_id(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))

        attachment = await attachment_service.save_attachment(
            db_session, "s1", "capture.png", "image/png", b"[FICTIF octets]"
        )

        disk_path = tmp_path / attachment.file_path
        assert disk_path.read_bytes() == b"[FICTIF octets]"
        assert attachment.id is not None  # généré avant l'écriture disque, pas au flush
        assert disk_path.name == f"{attachment.id}.png"  # jamais dérivé de `filename`
        assert attachment.filename == "capture.png"  # conservé en métadonnée pour l'affichage
        assert attachment.kind == "image"

    @pytest.mark.asyncio
    async def test_two_files_with_same_extension_do_not_overwrite_each_other(
        self, db_session, monkeypatch, tmp_path
    ):
        """Régression : l'id n'était généré qu'au flush, donc tous les fichiers
        d'une session s'appelaient "None.<ext>" et le second écrasait le premier."""
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))

        first = await attachment_service.save_attachment(db_session, "s1", "a.pdf", "application/pdf", b"[FICTIF] A")
        second = await attachment_service.save_attachment(db_session, "s1", "b.pdf", "application/pdf", b"[FICTIF] B")

        assert first.file_path != second.file_path
        assert (tmp_path / first.file_path).read_bytes() == b"[FICTIF] A"
        assert (tmp_path / second.file_path).read_bytes() == b"[FICTIF] B"
        assert "None" not in first.file_path

    @pytest.mark.asyncio
    async def test_path_traversal_filename_does_not_escape_upload_dir(self, db_session, monkeypatch, tmp_path):
        """Un nom de fichier malveillant ne doit jamais influencer le chemin
        disque réel (seul l'id, un UUID généré côté serveur, le fait)."""
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))

        attachment = await attachment_service.save_attachment(
            db_session, "s1", "../../../../etc/passwd", "text/plain", b"contenu"
        )

        disk_path = tmp_path / attachment.file_path
        assert disk_path.is_relative_to(tmp_path)
        assert disk_path.parent == tmp_path / "s1"


class TestProcessAttachmentDocuments:
    @pytest.mark.asyncio
    async def test_extracts_txt(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "notes.txt", "text/plain", "[FICTIF] besoin urgent".encode("utf-8")
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert attachment.extracted_text == "[FICTIF] besoin urgent"
        assert attachment.extraction_error is None

    @pytest.mark.asyncio
    async def test_extracts_pdf(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        pdf_bytes = _build_minimal_pdf("[FICTIF] Capture PDF")
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "spec.pdf", "application/pdf", pdf_bytes
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert "[FICTIF] Capture PDF" in attachment.extracted_text
        assert attachment.extraction_error is None

    @pytest.mark.asyncio
    async def test_extracts_docx(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        docx_bytes = _build_docx("[FICTIF] Paragraphe Word")
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "cahier.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx_bytes,
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert "[FICTIF] Paragraphe Word" in attachment.extracted_text
        assert attachment.extraction_error is None

    @pytest.mark.asyncio
    async def test_extracts_xlsx(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        xlsx_bytes = _build_xlsx("[FICTIF] Cellule Excel")
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "donnees.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_bytes,
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert "[FICTIF] Cellule Excel" in attachment.extracted_text
        assert attachment.extraction_error is None

    @pytest.mark.asyncio
    async def test_extracts_pptx(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        pptx_bytes = _build_pptx("[FICTIF] Diapositive PowerPoint")
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "presentation.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            pptx_bytes,
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert "[FICTIF] Diapositive PowerPoint" in attachment.extracted_text
        assert attachment.extraction_error is None

    @pytest.mark.asyncio
    async def test_corrupted_pdf_sets_extraction_error_not_exception(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "corrompu.pdf", "application/pdf", b"ceci n'est pas un vrai PDF"
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert attachment.extracted_text is None
        assert attachment.extraction_error is not None

    @pytest.mark.asyncio
    async def test_unsupported_extension_sets_extraction_error(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "archive.zip", "application/zip", b"PK\x03\x04"
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert attachment.extracted_text is None
        assert "non pris en charge" in attachment.extraction_error

    @pytest.mark.asyncio
    async def test_missing_file_on_disk_sets_extraction_error(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        attachment = AttachmentModel(
            session_id="s1", filename="fantome.txt", content_type="text/plain",
            kind="document", file_path="s1/does-not-exist.txt",
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert attachment.extraction_error is not None


class TestProcessAttachmentImages:
    @pytest.mark.asyncio
    async def test_describes_image_via_vision_provider(self, db_session, monkeypatch, tmp_path):
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "capture.png", "image/png", b"[FICTIF octets image]"
        )

        await attachment_service.process_attachment(attachment, _FixedVisionProvider())

        assert attachment.extracted_text == "[FICTIF] Description factice de l'image."
        assert attachment.extraction_error is None

    @pytest.mark.asyncio
    async def test_vision_failure_sets_extraction_error_not_exception(self, db_session, monkeypatch, tmp_path):
        """Best-effort : Ollama down / modèle vision absent ne doit jamais
        faire planter le traitement de la pièce jointe."""
        monkeypatch.setattr(attachment_service.settings, "UPLOAD_DIR", str(tmp_path))
        attachment = await attachment_service.save_attachment(
            db_session, "s1", "capture.png", "image/png", b"[FICTIF octets image]"
        )

        await attachment_service.process_attachment(attachment, _FailingVisionProvider())

        assert attachment.extracted_text is None
        assert attachment.extraction_error == "[failing] modèle vision indisponible"


class TestBuildAttachmentsContext:
    def test_returns_none_when_nothing_extracted(self):
        attachments = [
            AttachmentModel(session_id="s1", filename="a.txt", content_type="text/plain", kind="document"),
        ]
        assert attachment_service.build_attachments_context(attachments) is None

    def test_concatenates_extracted_content_with_filenames(self):
        a = AttachmentModel(session_id="s1", filename="a.txt", content_type="text/plain", kind="document")
        a.extracted_text = "[FICTIF] contenu A"
        b = AttachmentModel(session_id="s1", filename="b.png", content_type="image/png", kind="image")
        b.extracted_text = "[FICTIF] description B"

        result = attachment_service.build_attachments_context([a, b])

        assert "a.txt" in result
        assert "[FICTIF] contenu A" in result
        assert "b.png" in result
        assert "[FICTIF] description B" in result

    def test_ignores_attachments_without_extracted_text(self):
        failed = AttachmentModel(session_id="s1", filename="bad.zip", content_type="application/zip", kind="document")
        failed.extraction_error = "Type non pris en charge"

        assert attachment_service.build_attachments_context([failed]) is None
